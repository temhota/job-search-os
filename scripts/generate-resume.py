#!/usr/bin/env python3
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_font(run, name="Arial", size=10.2, bold=False, color="20212A"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=60, start=0, bottom=60, end=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_link(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2D63B8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    set_font(r, size=11.2, bold=True, color="000000")
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "2D63B8")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), size=9.7)


def build(payload, output_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(2)

    title_style = doc.styles["Title"]
    title_style.font.color.rgb = RGBColor.from_string("000000")
    style_pr = title_style.element.get_or_add_pPr()
    style_border = style_pr.find(qn("w:pBdr"))
    if style_border is not None:
        style_pr.remove(style_border)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run(payload["name"])
    set_font(run, size=24, bold=True, color="000000")
    title_pr = title._p.get_or_add_pPr()
    existing_border = title_pr.find(qn("w:pBdr"))
    if existing_border is not None:
        title_pr.remove(existing_border)
    no_border = OxmlElement("w:pBdr")
    no_bottom = OxmlElement("w:bottom")
    no_bottom.set(qn("w:val"), "nil")
    no_border.append(no_bottom)
    title_pr.append(no_border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(3)
    set_font(subtitle.add_run(payload["headline"]), size=12.8, bold=True, color="20212A")

    contact = doc.add_paragraph()
    contact.paragraph_format.space_after = Pt(4)
    for index, item in enumerate(payload["contactLine"].split(" | ")):
        if index:
            set_font(contact.add_run(" | "), size=9.5, color="555966")
        if "@" in item:
            add_link(contact, item, f"mailto:{item}")
        elif item.startswith(("https://", "http://")):
            add_link(contact, item, item)
        elif item.startswith(("linkedin.com/", "www.linkedin.com/")):
            add_link(contact, "LinkedIn", f"https://{item}")
        else:
            set_font(contact.add_run(item), size=9.5, color="555966")

    german = payload.get("language") == "German"
    section_heading(doc, "Profil" if german else "Professional Summary")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.03
    set_font(p.add_run(payload["summary"]), size=9.8)

    section_heading(doc, "Kernkompetenzen" if german else "Core Skills")
    for line in payload["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.2)
        label, _, rest = line.partition(":")
        set_font(p.add_run(label + (":" if rest else "")), size=9.6, bold=True, color="2D63B8")
        if rest:
            set_font(p.add_run(rest), size=9.6)

    section_heading(doc, "Berufserfahrung" if german else "Professional Experience")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(payload["careerNote"]), size=9.8, bold=True)

    for item in payload["experience"]:
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(5.7)
        table.columns[1].width = Inches(1.35)
        left, right = table.rows[0].cells
        set_cell_margins(left)
        set_cell_margins(right)
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        set_font(lp.add_run(f'{item["role"]} | {item["company"]}'), size=10.1, bold=True)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        set_font(rp.add_run(item["dates"]), size=9.5, bold=True, color="555966")
        for bullet in item["bullets"]:
            add_bullet(doc, bullet)

    section_heading(doc, "Ausbildung" if german else "Education")
    for item in payload["education"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(item), size=9.6, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("SPRACHEN: " if german else "LANGUAGES: "), size=9.6, bold=True, color="2D63B8")
    set_font(p.add_run(payload["languages"]), size=9.6)

    props = doc.core_properties
    props.title = f'{payload["name"]} {payload["headline"]}'
    props.subject = "Tailored resume"
    props.author = payload["name"]
    props.keywords = ", ".join(payload["skills"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: generate-resume.py payload.json output.docx")
    payload = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    build(payload, Path(sys.argv[2]))
